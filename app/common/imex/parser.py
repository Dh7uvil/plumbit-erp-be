"""Parse xlsx/csv workbooks and apply column mappings. No business rules."""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from typing import Any

from app.common.imex.schemas import ImexColumn, ImexField, ImexMappingEntry
from app.core.exceptions import ValidationError

MAX_IMPORT_ROWS = 2000


def parse_tabular(filename: str | None, content: bytes) -> tuple[list[str], list[list[str]]]:
    """Return header cells and data rows. Detects the first non-empty row as headers."""

    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = [[cell.strip() for cell in row] for row in reader]
    else:
        rows = _read_xlsx(content)
    header_index = next((index for index, row in enumerate(rows) if any(row)), None)
    if header_index is None:
        raise ValidationError("The file has no header row")
    headers = [cell or f"Column {index + 1}" for index, cell in enumerate(rows[header_index])]
    data = rows[header_index + 1 :]
    if len(data) > MAX_IMPORT_ROWS:
        raise ValidationError(
            "Import exceeds the maximum row count",
            details={"max_rows": MAX_IMPORT_ROWS, "row_count": len(data)},
        )
    return headers, data


def suggest_mapping(headers: Sequence[str], fields: Sequence[ImexField]) -> list[ImexMappingEntry]:
    by_alias: dict[str, str] = {}
    for field in fields:
        by_alias[_norm(field.name)] = field.name
        by_alias[_norm(field.label)] = field.name
        for alias in field.aliases:
            by_alias[_norm(alias)] = field.name
    mapping: list[ImexMappingEntry] = []
    used: set[str] = set()
    for header in headers:
        field = by_alias.get(_norm(header))
        if field is None or field in used:
            continue
        used.add(field)
        mapping.append(ImexMappingEntry(column=header, field=field))
    return mapping


def apply_mapping(
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    mapping: Sequence[ImexMappingEntry],
) -> list[dict[str, str]]:
    index_by_header = {header: index for index, header in enumerate(headers)}
    mapped: list[dict[str, str]] = []
    for row in rows:
        if not any(str(cell).strip() for cell in row):
            continue
        values: dict[str, str] = {}
        for entry in mapping:
            index = index_by_header.get(entry.column)
            if index is None or index >= len(row):
                continue
            values[entry.field] = str(row[index]).strip()
        mapped.append(values)
    return mapped


def columns_from_headers(headers: Sequence[str]) -> list[ImexColumn]:
    return [ImexColumn(index=index, header=header) for index, header in enumerate(headers)]


def sample_rows(
    headers: Sequence[str], rows: Sequence[Sequence[str]], *, limit: int = 5
) -> list[dict[str, Any]]:
    samples: list[dict[str, Any]] = []
    for row in rows[:limit]:
        samples.append(
            {
                header: row[index] if index < len(row) else ""
                for index, header in enumerate(headers)
            }
        )
    return samples


def write_xlsx(headers: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell).strip() for cell in row])
    return rows


def _norm(value: str) -> str:
    return "".join(ch for ch in value.lower() if ch.isalnum())
